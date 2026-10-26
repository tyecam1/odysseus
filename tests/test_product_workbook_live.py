"""Live-workbook regression tests (corrective round 2026-10-26).

The merged workbook wrote computed values, static per-decile fills and no
formulas. These tests require the repaired contract:

* Overview composite scores are LIVE Excel formulas over Details per-row
  factor and weight INPUT cells, so editing an input recalculates -
  including the missing-applicable penalty via COUNT().
* The score gradient is a genuine Excel conditional format (colour scale),
  not static per-cell fills, and Overview stays exactly 23 columns (A:W).
* Recalculated values (via LibreOffice, when installed) match the Python
  scoring engine.
"""
import subprocess
import shutil
from io import BytesIO

import pytest

from src.product_research import GAMMA, build_factors, composite_score
from src.product_workbook import (
    DETAIL_COLUMNS, FACTOR_INPUT_COLUMNS, OVERVIEW_COLUMNS, WEIGHT_COLUMNS,
    WEIGHT_DEFAULTS, build_workbook,
)

pytest.importorskip("openpyxl")

from openpyxl import load_workbook  # noqa: E402

GAMMA_DEFAULT = GAMMA


def rec(name, category="Cannabis flower", **kw):
    r = {"product_name": name, "category": category, "buying_status": "purchasable"}
    r.update(kw)
    return r


def _mixed_records():
    return [
        rec("FlowerA", combined_rating=4.5, unit_price=10.0, unit_basis="per g",
            leafly_thc=25, sativa_pct=75, indica_pct=25, leafly_rating_count=120,
            seller_popularity_rank=1, seller_tier="AA",
            overview_options=[{"size": "1g", "price": 10.0}]),
        rec("FlowerB", combined_rating=3.5, unit_price=20.0, unit_basis="per g",
            leafly_thc=18, sativa_pct=40, indica_pct=60, leafly_rating_count=30,
            seller_popularity_rank=3, seller_tier="A"),
        rec("EdibleX", category="Edibles", combined_rating=4.0, unit_price=5.0,
            unit_basis="per item", leafly_rating_count=10, seller_popularity_rank=2),
    ]


@pytest.fixture(scope="module")
def workbook_blob():
    scored = build_factors(_mixed_records())
    return build_workbook(scored, "2026-10-26"), scored


def test_overview_scores_are_live_formulas(workbook_blob):
    blob, _ = workbook_blob
    wb = load_workbook(BytesIO(blob))
    ws = wb["Overview"]
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 4).value
        assert isinstance(value, str) and value.startswith("="), \
            f"Overview D{row} must be a live formula, got {value!r}"
        assert "Details!" in value and "COUNT(" in value


def test_overview_stays_exactly_23_columns(workbook_blob):
    blob, _ = workbook_blob
    wb = load_workbook(BytesIO(blob))
    ws = wb["Overview"]
    assert ws.max_column == 23  # canon contract: exact A:W
    assert [c.value for c in ws[1]] == OVERVIEW_COLUMNS


def test_details_factor_inputs_and_calculation_chain_are_present(workbook_blob):
    blob, _ = workbook_blob
    wb = load_workbook(BytesIO(blob))
    ws = wb["Details"]
    headers = [c.value for c in ws[1]]
    for name in FACTOR_INPUT_COLUMNS + WEIGHT_COLUMNS:
        assert name in headers
    for name in ("GM (live)", "B (live)", "C (live)"):
        col = headers.index(name) + 1
        value = ws.cell(2, col).value
        assert isinstance(value, str) and value.startswith("="), \
            f"Details {name} must be a live formula, got {value!r}"


def test_weight_and_gamma_input_cells_carry_canon_defaults(workbook_blob):
    blob, _ = workbook_blob
    wb = load_workbook(BytesIO(blob))
    ws = wb["Details"]
    headers = [c.value for c in ws[1]]
    for name, default in WEIGHT_DEFAULTS.items():
        header = ("Weight: sativa/indica (input)" if name == "si"
                  else f"Weight: {name} (input)")
        col = headers.index(header) + 1
        assert ws.cell(2, col).value == pytest.approx(default)
    assert ws.cell(2, headers.index("Gamma (input)") + 1).value == GAMMA_DEFAULT


def test_gradient_is_conditional_format_not_static_fills(workbook_blob):
    blob, _ = workbook_blob
    wb = load_workbook(BytesIO(blob))
    ws = wb["Overview"]
    rules = [r for rng in ws.conditional_formatting for r in rng.rules
             if "D" in str(rng.sqref)]
    assert any(r.type == "colorScale" for r in rules), \
        "score column must carry a genuine Excel colour-scale conditional format"
    fills = {ws.cell(row, 4).fill.fgColor.rgb for row in range(2, ws.max_row + 1)}
    assert len(fills) == 1, f"static decile fills are gone; saw {fills}"


def _recalculated(blob, tmp_path):
    """LibreOffice-recalculated copy; returns (overview_values, details_values)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:  # pragma: no cover - LibreOffice missing on the host
        pytest.skip("LibreOffice not installed")
    src = tmp_path / "in.xlsx"
    src.write_bytes(blob)
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    profile = tmp_path / "lo-profile"
    profile.mkdir(exist_ok=True)
    subprocess.run(
        [soffice, "-env:UserInstallation=file://" + str(profile), "--headless",
         "--convert-to", "xlsx", "--outdir", str(out), str(src)],
        check=True, capture_output=True, timeout=180)
    produced = out / "in.xlsx"
    assert produced.exists()
    wb = load_workbook(produced, data_only=True)
    return wb["Overview"], wb["Details"]


def test_recalculated_scores_match_python_engine(workbook_blob, tmp_path):
    blob, scored = workbook_blob
    overview, _ = _recalculated(blob, tmp_path)
    expected = sorted(round(s["percent"], 6) for _, _, s in scored)
    got = sorted(round(overview.cell(row, 4).value * 100, 4)
                 for row in range(2, overview.max_row + 1))
    assert len(got) == len(expected)
    for g, e in zip(got, expected):
        assert g == pytest.approx(e, abs=0.02), (g, e)


def test_editing_a_factor_input_recalculates_the_score(workbook_blob, tmp_path):
    blob, scored = workbook_blob
    wb_edit = load_workbook(BytesIO(blob))
    details = wb_edit["Details"]
    headers = [c.value for c in details[1]]
    thc_col = headers.index(FACTOR_INPUT_COLUMNS[3]) + 1  # THC factor input
    original = details.cell(2, thc_col).value
    assert original and original != 1
    details.cell(2, thc_col).value = min(1.5, original * 1.2)
    edited = BytesIO()
    wb_edit.save(edited)
    overview, _ = _recalculated(edited.getvalue(), tmp_path)
    new_score = overview.cell(2, 4).value * 100
    _, _, s0 = scored[0]
    assert new_score - s0["percent"] > 0.1  # a bigger THC advantage raises the score


def test_editing_weights_recalculates_all_scores(workbook_blob, tmp_path):
    blob, scored = workbook_blob
    wb_edit = load_workbook(BytesIO(blob))
    ws = wb_edit["Details"]
    headers = [c.value for c in ws[1]]
    pop_col = headers.index("Weight: popularity (input)") + 1
    for row in range(2, ws.max_row + 1):
        ws.cell(row, pop_col).value = 0.0  # popularity removed from the model
    edited = BytesIO()
    wb_edit.save(edited)
    overview, _ = _recalculated(edited.getvalue(), tmp_path)
    for row, (rec_, factors, score) in zip(range(2, overview.max_row + 1), scored):
        expected = composite_score(factors, rec_["category"],
                                   weight_overrides={"popularity": 0.0})["percent"]
        got = overview.cell(row, 4).value * 100
        assert got == pytest.approx(expected, abs=0.02), \
            (rec_["product_name"], got, expected)


def test_clearing_factor_inputs_moves_the_missing_penalty(workbook_blob, tmp_path):
    """The missing-applicable count is live: clearing two present Edibles
    factors must apply the 0.8^max(0,m-1) penalty exactly as the engine."""
    blob, scored = workbook_blob
    target = next((rec_, f, s) for rec_, f, s in scored
                  if rec_["category"] == "Edibles" and len(s["factors_present"]) >= 3)
    row = [i for i, (rec_, _, _) in enumerate(scored, start=2)
           if rec_ is target[0]][0]
    wb_edit = load_workbook(BytesIO(blob))
    ws = wb_edit["Details"]
    headers = [c.value for c in ws[1]]
    # Edibles: rating/price/popularity applicable; clear rating and price
    for factor in ("Rating factor (input)", "Price factor (input)"):
        ws.cell(row, headers.index(factor) + 1).value = None
    edited = BytesIO()
    wb_edit.save(edited)
    overview, _ = _recalculated(edited.getvalue(), tmp_path)
    got = overview.cell(row, 4).value * 100
    expected = composite_score(target[1], "Edibles")["percent"] \
        if False else _engine_percent_cleared(target[1], "Edibles")
    assert got == pytest.approx(expected, abs=0.02), (got, expected)


def _engine_percent_cleared(factors, category):
    """Engine score with rating and price factors cleared (rebuilt Factors)."""
    from src.product_research import Factors
    cleared = Factors(popularity=factors.popularity)
    cleared.detail = factors.detail
    return composite_score(cleared, category)["percent"]
