"""Tests for the Green Team x Leafly scoring engine and workbook contract."""
import math
import copy
import json
import re
from io import BytesIO

from openpyxl import load_workbook

import pytest

from src.product_research import (
    APPLICABLE, GAMMA, _midrank_factor, apply_rarity_grade, build_factors,
    combined_popularity, combined_rating, composite_score, parse_quantity_grams,
    seller_grade_numeric, sort_records, unit_price,
)
from src.product_workbook import DETAIL_COLUMNS, OVERVIEW_COLUMNS, build_workbook


def _load(blob):
    """Read the generated artifact with the supported openpyxl reader."""
    return load_workbook(BytesIO(blob))


def rec(name, category="Cannabis flower", **kw):
    r = {"product_name": name, "category": category, "buying_status": "purchasable"}
    r.update(kw)
    return r


# --- unit prices -------------------------------------------------------------

def test_unit_price_min_per_gram_excludes_zero_and_unknown():
    opts = [{"size": "1g", "price": 12.5}, {"size": "3.5g", "price": 30},
            {"size": "7g", "price": 0}, {"size": "1 PACK", "price": 20},
            {"size": "?", "price": 5}]
    best, basis, det = unit_price(opts)
    assert best == pytest.approx(30 / 3.5)
    assert basis == "per g"
    assert len(det["excluded_options"]) == 3


def test_parse_quantity_units():
    assert parse_quantity_grams("3.5g") == 3.5
    assert parse_quantity_grams("1000MG") == 1.0
    assert parse_quantity_grams("1 PACK") is None
    assert parse_quantity_grams("") is None


# --- ratings -----------------------------------------------------------------

def test_unrated_is_not_zero_stars():
    blended, detail = combined_rating((None, None), (0.0, 0))
    assert blended is None
    blended2, _ = combined_rating((None, None), (None, None))
    assert blended2 is None


def test_combined_rating_sample_weighting():
    # leafly 4.5 with 100 counts dominates seller 3.0 with 1 count
    blended, detail = combined_rating((4.5, 100), (3.0, 1))
    assert blended > 3.0 and blended < 4.5
    assert (blended - 3.0) > (4.5 - blended)  # pulled toward leafly
    assert detail["weights"]["leafly"] > detail["weights"]["seller"]


def test_combined_popularity_uses_verified_seller_ordinal():
    f, det = combined_popularity(1000, 1, 10, "Cannabis flower",
                                 [1000, 10, 5], [1, 2, 3])
    assert f is not None and det["seller_factor"] is not None
    # rank 1 (most popular) must score above rank 3
    f3, _ = combined_popularity(1000, 3, 10, "Cannabis flower",
                                [1000, 10, 5], [1, 2, 3])
    assert f > f3


def test_zero_reviews_is_not_zero_demand():
    f, _ = combined_popularity(None, 1, 10, "Hash and sift", [None, None], [1, 2])
    assert f is not None  # seller ordinal still carries popularity evidence


# --- scoring -----------------------------------------------------------------

def test_gamma_spread_midpoint_and_range():
    f = dict(grade=(1.0, 1.0))
    mid = composite_score(type("F", (), {"grade": (1.0, 1.0), "rating": 1.0,
                                          "price": (1.0, "per g"), "thc": 1.0,
                                          "si": 1.0, "popularity": 1.0,
                                          "detail": {}})(), "Cannabis flower")
    assert mid["percent"] == pytest.approx(50.0)
    # strong advantage approaches (but never fabricates) 100
    strong = composite_score(type("F", (), {"grade": (1.5, 1.0), "rating": 1.5,
                                             "price": (1.5, "per g"), "thc": 1.5,
                                             "si": 1.5, "popularity": 1.5,
                                             "detail": {}})(), "Cannabis flower")
    assert strong["percent"] > 90
    assert strong["percent"] <= 100


def test_missing_penalty_first_forgiven_then_compounded():
    # Cannabis flower: grade present only -> 5 applicable missing
    f1 = type("F", (), {"grade": (1.0, 1.0), "rating": None, "price": None,
                        "thc": None, "si": None, "popularity": None,
                        "detail": {}})()
    r1 = composite_score(f1, "Cannabis flower")
    assert r1["missing_penalty"] == pytest.approx(0.8 ** 4)
    # all applicable present -> penalty 1
    tall = type("F", (), {"grade": (1.0, 1.0), "rating": 1.0, "price": (1.0, "g"),
                          "thc": 1.0, "si": 1.0, "popularity": 1.0, "detail": {}})()
    assert composite_score(tall, "Cannabis flower")["missing_penalty"] == 1.0


def test_inapplicable_not_penalised_mushrooms():
    # mushroom with rating+price+popularity present => no missing metrics
    f = type("F", (), {"grade": None, "rating": 1.0, "price": (1.0, "per g"),
                       "thc": None, "si": None, "popularity": 1.0, "detail": {}})()
    r = composite_score(f, "Mushroom-related")
    assert r["missing_applicable"] == []
    assert r["missing_penalty"] == 1.0


def test_popularity_weight_actually_25_percent():
    # with everything else equal, popularity alone cannot move the score much
    base = type("F", (), {"grade": (1.2, 1.0), "rating": 1.2, "price": (1.2, "g"),
                          "thc": 1.2, "si": 1.2, "popularity": 1.0, "detail": {}})()
    pop = type("F", (), {"grade": (1.2, 1.0), "rating": 1.2, "price": (1.2, "g"),
                         "thc": 1.2, "si": 1.2, "popularity": 1.5, "detail": {}})()
    a = composite_score(base, "Cannabis flower")["percent"]
    b = composite_score(pop, "Cannabis flower")["percent"]
    assert 0 < b - a < 3.0  # weight 0.25 keeps the effect small


def test_single_peer_category_is_neutral():
    recs = [rec("Solo", "Vapes", combined_rating=4.0, unit_price=20.0,
                unit_basis="per g", leafly_rating_count=10,
                seller_popularity_rank=1)]
    scored = build_factors(recs)
    _, f, score = scored[0]
    assert f.rating == 1.0 and f.price[0] == 1.0
    assert score["gm"] == 1.0


def test_seller_only_singleton_popularity_is_present_and_neutral():
    _, factors, score = build_factors([rec("Synthetic", "Vapes", seller_popularity_rank=1)])[0]
    assert factors.popularity == 1.0
    assert "popularity" not in score["missing_applicable"]


def test_present_flower_reference_thc_and_orientation_are_scored():
    record = rec("Synthetic flower", leafly_thc=25, sativa_pct=75, indica_pct=25,
                 primary_leafly_source="https://example.invalid/cultivar")
    before = copy.deepcopy(record)
    _, factors, score = build_factors([record])[0]
    assert factors.thc == factors.si == 1.0
    assert {"thc", "si"}.isdisjoint(score["missing_applicable"])
    assert record == before
    assert factors.detail["thc"]["evidence_class"] == "cultivar_reference"
    assert "never a batch measurement" in factors.detail["thc"]["basis"]
    assert factors.detail["thc"]["source"] == record["primary_leafly_source"]
    assert factors.detail["si"]["raw_values"] == {"sativa_pct": 75, "indica_pct": 25}


def test_flower_reference_peers_ties_ordering_and_mean():
    records = [rec(f"Synthetic {i}", leafly_thc=thc, sativa_pct=sativa, indica_pct=100-sativa)
               for i, (thc, sativa) in enumerate([(10, 20), (25, 75), (25, 75), (30, 90)])]
    scored = build_factors(records)
    for metric in ("thc", "si"):
        values = [getattr(f, metric) for _, f, _ in scored]
        assert values[0] < values[1] == values[2] < values[3]
        assert sum(values) / len(values) == pytest.approx(1)
    scores = [score["score"] for _, _, score in scored]
    assert scores[0] < scores[1] == scores[2] < scores[3]
    equal = build_factors([rec(f"Synthetic equal {i}", leafly_thc=25, sativa_pct=75)
                           for i in range(3)])
    assert all(f.thc == f.si == 1 for _, f, _ in equal)


@pytest.mark.parametrize("category", sorted(set(APPLICABLE) - {"Cannabis flower"}))
def test_nonflower_reference_values_never_become_factors_or_peers(category):
    flower = rec("Synthetic flower", leafly_thc=25, sativa_pct=75)
    other = rec("Synthetic nonflower", category, leafly_thc=100, sativa_pct=100, indica_pct=0)
    scored = {r["product_name"]: (f, score) for r, f, score in build_factors([flower, other])}
    assert scored["Synthetic flower"][0].thc == scored["Synthetic flower"][0].si == 1
    factors, score = scored["Synthetic nonflower"]
    assert factors.thc is factors.si is None
    assert {"thc", "si"}.isdisjoint(score["missing_applicable"] + score["factors_present"])


@pytest.mark.parametrize("value", [None, True, "25", -1, 101, float("nan"), float("inf")])
def test_invalid_flower_reference_values_are_missing_and_excluded_from_peers(value):
    records = [rec("Synthetic valid", leafly_thc=25, sativa_pct=75),
               rec("Synthetic invalid", leafly_thc=value, sativa_pct=value)]
    (_, valid, _), (_, invalid, score) = build_factors(records)
    assert valid.thc == valid.si == 1
    assert invalid.thc is invalid.si is None
    assert {"thc", "si"} <= set(score["missing_applicable"])
    assert invalid.detail["thc"]["excluded_reason"]
    assert invalid.detail["si"]["excluded_reason"]


def test_orientation_single_axis_comparison_keeps_source_missingness():
    records = [rec("Synthetic sativa", sativa_pct=75),
               rec("Synthetic indica", indica_pct=25),
               rec("Synthetic conflicting", sativa_pct=75, indica_pct=75)]
    before = copy.deepcopy(records)
    (_, sativa, _), (_, indica, _), (_, conflicting, score) = build_factors(records)
    assert sativa.si == indica.si == 1
    assert indica.detail["si"]["comparison_value"] == 75
    assert indica.detail["si"]["raw_values"] == {"sativa_pct": None, "indica_pct": 25}
    assert conflicting.si is None and "si" in score["missing_applicable"]
    assert records == before


def test_rarity_grade_only_for_ungraded_flower():
    recs = [rec("Synthetic ungraded"), rec("Synthetic seller", seller_tier="A"),
            rec("Synthetic edible", category="Edibles")]
    policy = {r["product_name"]: {"tier": "AA", "confidence": 0.6,
              "provenance": {"source": "synthetic-policy", "revision": "test-v1"}}
              for r in recs}
    apply_rarity_grade(recs)
    assert all("inferred_grade" not in r for r in recs)  # no bundled personal policy
    build_factors(recs, rarity_policy=policy)
    assert recs[0]["inferred_grade"]["tier"] == "AA"
    assert recs[0]["inferred_grade"]["numeric"] == 2.0
    assert recs[0]["inferred_grade"]["provenance"] == policy["Synthetic ungraded"]["provenance"]
    assert "inferred_grade" not in recs[1]  # seller grade preserved, no inference
    assert "inferred_grade" not in recs[2]  # not a flower


def test_reported_grade_overrides_stale_inference_on_repeated_build():
    shared = dict(combined_rating=4.0, unit_price=10.0, unit_basis="per g",
                  leafly_thc=25, sativa_pct=75, leafly_rating_count=10,
                  seller_popularity_rank=1)
    record = rec("Synthetic updated", **shared)
    peer = rec("Synthetic peer", seller_tier="AA", **shared)
    unrelated = [rec("Synthetic vape", "Vapes", combined_rating=4.5),
                 rec("Synthetic vape peer", "Vapes", combined_rating=3.5)]
    records = [record, peer, *unrelated]
    fresh_records = copy.deepcopy(records)
    fresh_records[0]["seller_tier"] = "AAA"
    policy = {record["product_name"]: {
        "tier": "A", "confidence": 0.8, "basis": "synthetic rarity estimate",
        "provenance": {"source": "synthetic-policy", "revision": "test-v1"},
    }}

    first = build_factors(records, rarity_policy=policy)
    audit = copy.deepcopy(record["inferred_grade"])
    assert audit == {**policy[record["product_name"]], "numeric": 1.0}
    assert first[0][1].grade == pytest.approx((1 - 0.8 * 0.5 / 7, 0.8))
    assert first[0][2]["score"] < first[1][2]["score"]

    record["seller_tier"] = "AAA"
    rebuilt = build_factors(records, rarity_policy=policy)
    fresh = build_factors(fresh_records, rarity_policy=policy)
    assert rebuilt[0][0] is record
    assert record["inferred_grade"] == audit  # audit survives seller precedence
    assert record["seller_tier"] == "AAA"
    assert "inferred_grade" not in fresh_records[0]
    assert rebuilt[0][1].grade == pytest.approx((1 + 0.5 / 7, 1.0))
    assert rebuilt[1][1].grade == pytest.approx((1 - 0.5 / 7, 1.0))
    assert rebuilt[0][2]["score"] > rebuilt[1][2]["score"]
    assert [(f, score) for _, f, score in rebuilt] == [(f, score) for _, f, score in fresh]
    assert [r["product_name"] for r, _, _ in sort_records(rebuilt)] == [
        r["product_name"] for r, _, _ in sort_records(fresh)]
    for (_, before, _), (_, after, _) in zip(first, rebuilt):
        assert (before.rating, before.price, before.thc, before.si,
                before.popularity, before.detail) == (
                    after.rating, after.price, after.thc, after.si,
                    after.popularity, after.detail)
    assert rebuilt[2:] == first[2:]  # other-category scores are unchanged


@pytest.mark.parametrize("seller_tier", [None, "", "unknown"])
def test_inferred_grade_fallback_survives_repeated_build_without_valid_seller_grade(seller_tier):
    record = rec("Synthetic inferred", seller_tier=seller_tier)
    peer = rec("Synthetic peer", seller_tier="AA")
    policy = {record["product_name"]: {
        "tier": "A", "confidence": 0.8, "provenance": "synthetic-policy",
    }}
    first = build_factors([record, peer], rarity_policy=policy)
    before = copy.deepcopy(record)
    rebuilt = build_factors([record, peer])
    assert record == before
    assert rebuilt == first
    assert rebuilt[0][1].grade == pytest.approx((1 - 0.8 * 0.5 / 7, 0.8))
    assert rebuilt[0][2]["score"] < rebuilt[1][2]["score"]


@pytest.mark.parametrize("entry", [
    {"tier": "AA", "confidence": 0.5},
    {"tier": "unknown", "confidence": 0.5, "provenance": "synthetic"},
    {"tier": "AA", "confidence": -0.1, "provenance": "synthetic"},
    {"tier": "AA", "confidence": float("nan"), "provenance": "synthetic"},
])
def test_invalid_injected_grade_policy_rejected(entry):
    with pytest.raises(ValueError):
        apply_rarity_grade([rec("Synthetic")], {"Synthetic": entry})


@pytest.mark.parametrize("values", [[4], [4, 4], [4, 4, 4], [1, 2, 2, 4, 9]])
@pytest.mark.parametrize("smaller", [False, True])
def test_midrank_ties_mean_and_direction(values, smaller):
    factors = [_midrank_factor(values, i, 0.8, smaller) for i in range(len(values))]
    assert sum(factors) / len(factors) == pytest.approx(1)
    if len(set(values)) == 1:
        assert factors == [1.0] * len(values)
    else:
        ordered = sorted(zip(values, factors))
        for (a, fa), (b, fb) in zip(ordered, ordered[1:]):
            if a == b:
                assert fa == fb
            else:
                assert (fa > fb) if smaller else (fa < fb)


def test_price_peers_require_same_category_and_unit_basis():
    records = [rec("gram", "Vapes", unit_price=1, unit_basis="per g"),
               rec("item", "Vapes", unit_price=100, unit_basis="per item"),
               rec("other category", "Edibles", unit_price=10, unit_basis="per g")]
    assert all(f.price[0] == 1 for _, f, _ in build_factors(records))
    records.append(rec("expensive gram", "Vapes", unit_price=2, unit_basis="per g"))
    factors = {r["product_name"]: f for r, f, _ in build_factors(records)}
    assert factors["gram"].price[0] > 1 > factors["expensive gram"].price[0]
    assert factors["item"].price == (1.0, "per item")


@pytest.mark.parametrize("price,basis", [(0, "per g"), (-1, "per g"),
    (float("nan"), "per g"), (float("inf"), "per g"), ("1", "per g"),
    (True, "per g"), (1, None), (1, "")])
def test_invalid_price_inputs_have_reasons(price, basis):
    record = rec("Synthetic", "Vapes", unit_price=price, unit_basis=basis)
    _, factor, score = build_factors([record])[0]
    assert factor.price is None and "price" in score["missing_applicable"]
    assert factor.detail["price"]["excluded_reason"]


def test_invalid_options_excluded_with_source_quantities_preserved():
    options = [{"size": "2g", "price": 10, "variant_id": "synthetic-v1"},
               {"size": "1g", "price": "bad"}, {"size": "0g", "price": 1},
               {"size": "-1g", "price": 1}, {"size": "1g", "price": float("inf")},
               {"size": "1g", "price": True}, None]
    best, basis, detail = unit_price(options)
    assert (best, basis) == (5, "per g")
    assert len(detail["excluded_options"]) == 6
    assert all(detail["excluded_options"])
    assert detail["source_options"] == options
    assert detail["valid_options"] == [{"source_option_index": 0, "quantity_grams": 2,
                                        "unit_price": 5, "unit_basis": "per g"}]


def test_every_item_gets_a_score():
    recs = [
        rec("A", combined_rating=4.0, unit_price=10.0, unit_basis="per g",
            leafly_rating_count=50, seller_popularity_rank=1),
        rec("B", category="Edibles", combined_rating=None, unit_price=None),
        rec("C", category="Hash and sift", seller_tier="AA", combined_rating=3.0,
            unit_price=6.7, unit_basis="per g", leafly_rating_count=5,
            seller_popularity_rank=3),
    ]
    scored = build_factors(recs)
    assert all(isinstance(s["score"], float) and 0.0 <= s["score"] <= 1.0
               for _, _, s in scored)


# --- sorting / workbook ------------------------------------------------------

def _scored(recs):
    return build_factors(recs)


def test_availability_first_then_category_score_price():
    recs = [
        rec("CheapLow", combined_rating=3.0, unit_price=5.0, unit_basis="per g",
            leafly_rating_count=5, seller_popularity_rank=2),
        rec("ExpHigh", combined_rating=4.9, unit_price=50.0, unit_basis="per g",
            leafly_rating_count=50, seller_popularity_rank=1),
        rec("OffList", buying_status="not listed", combined_rating=4.0,
            unit_price=1.0, unit_basis="per g", leafly_rating_count=99,
            seller_popularity_rank=1),
    ]
    ordered = sort_records(_scored(recs))
    names = [r["product_name"] for r, _, _ in ordered]
    # purchasable first, higher score first; not-listed last despite best metrics
    assert names == ["ExpHigh", "CheapLow", "OffList"]


def test_workbook_contract_two_sheets_exact_columns():
    recs = [
        rec("A", combined_rating=4.0, unit_price=10.0, unit_basis="per g",
            leafly_rating_count=50, seller_popularity_rank=1,
            store_description="seller text", permalink="https://x/1",
            overview_options=[{"size": "1g", "price": 10.0}],
            leafly_description="leafly text", match_classification="exact"),
        rec("B", category="Edibles", buying_status="not listed"),
    ]
    wb = _load(build_workbook(sort_records(_scored(recs)), "2026-10-25"))
    assert wb.sheetnames == ["Overview", "Details"]
    hdr = [c.value for c in wb["Overview"][1]][:23]  # table part; weights sit beyond
    assert hdr == OVERVIEW_COLUMNS and len(hdr) == 23
    hdr2 = [c.value for c in wb["Details"][1]]
    n_parts = len(hdr2) - len(DETAIL_COLUMNS)
    assert n_parts >= 1  # raw JSON part 1 always has a column
    assert hdr2 == DETAIL_COLUMNS + [f"Raw record JSON part {i}"
                                     for i in range(1, n_parts + 1)]


def test_workbook_losslessness_and_alignment():
    recs = [
        rec("A", combined_rating=4.0, unit_price=10.0, unit_basis="per g",
            leafly_rating_count=50, seller_popularity_rank=1,
            store_description="FULL seller description " * 10,
            permalink="https://seller/a", overview_options=[{"size": "1g", "price": 10.0}],
            leafly_description="leafly text", match_classification="exact"),
    ]
    wb = _load(build_workbook(sort_records(_scored(recs)), "2026-10-25"))
    assert wb["Overview"].max_row == wb["Details"].max_row  # row-aligned
    text = " ".join(str(c.value) for row in wb["Overview"].iter_rows() for c in row)
    assert ("FULL seller description " * 10).strip() in text
    d2 = " ".join(str(c.value) for row in wb["Details"].iter_rows() for c in row)
    assert "https://seller/a" in d2 and "leafly text" in d2


def test_workbook_row_order_matches_sorted_scores():
    recs = [
        rec("Low", combined_rating=3.0, unit_price=5.0, unit_basis="per g",
            leafly_rating_count=5, seller_popularity_rank=2),
        rec("High", combined_rating=4.9, unit_price=50.0, unit_basis="per g",
            leafly_rating_count=50, seller_popularity_rank=1),
    ]
    wb = _load(build_workbook(sort_records(_scored(recs)), "2026-10-25"))
    assert wb["Overview"].cell(2, 2).value == "High"


def test_workbook_formats_widths_and_conditional_score_gradient():
    record = rec("Synthetic", unit_price=5, unit_basis="per g", leafly_thc=25,
                 sativa_pct=75, indica_pct=25, overview_options=[{"size": "2g", "price": 10}])
    scored = build_factors([record])
    rows = [(dict(record, product_name=f"Synthetic {i}"), scored[0][1],
             dict(scored[0][2], score=i / 10, percent=i * 10)) for i in range(11)]
    wb = _load(build_workbook(rows, "2026-10-25"))
    ws = wb["Overview"]
    # the score gradient is a genuine Excel colour scale on the score column only
    ranges = list(ws.conditional_formatting)
    rules = [r for rng in ranges for r in rng.rules]
    assert len(rules) == 1 and rules[0].type == "colorScale"
    assert str(ranges[0].sqref) == "D2:D12"
    colors = [c.rgb for c in rules[0].colorScale.color]
    assert len(set(colors)) == 3 and all(re.fullmatch("00[0-9A-F]{6}", c) for c in colors)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column == 4:
                assert cell.number_format == "0.00%"
                assert isinstance(cell.value, str) and cell.value.startswith("=")
            elif cell.column == 7:
                assert cell.number_format == '"£"0.00\\ \\p\\e\\r\\ \\g'
            elif cell.column in (8, 9, 10):
                assert cell.number_format == '0.0"%"'
            elif cell.column >= 12 and cell.column % 2 == 0 and cell.value is not None:
                assert cell.number_format == '"£"0.00'
        assert ws.cell(row[0].row, 8).value == 25  # source percentage kept as points
    assert ws.freeze_panes == "A2" and wb["Details"].freeze_panes == "A2"
    for sheet in (ws, wb["Details"]):
        widths = [d.width for d in sheet.column_dimensions.values()]
        assert len(widths) >= 23 and all(w >= 14 for w in widths)


def test_workbook_unit_price_basis_is_visible_literal_and_totals_stay_currency():
    bases = ["per g", "per item", "per g", None, "", "unknown",
             'per "item";[Red]0\\<&>', "per \x00;[Red]0"]
    records = [rec(f"Synthetic {i}", unit_price=5 + i / 10, unit_basis=basis,
                   overview_options=[{"size": f"Option {k}", "price": 10 + k} for k in range(6)])
               for i, basis in enumerate(bases)]
    records.append(rec("No supplied basis", unit_price=8,
                       overview_options=[{"size": "1g", "price": 8}]))
    before = copy.deepcopy(records)
    wb = _load(build_workbook(build_factors(records), "2026-10-25"))
    ws, details = wb["Overview"], wb["Details"]
    assert [c.value for c in ws[1]][:23] == OVERVIEW_COLUMNS
    assert ws.max_row == details.max_row == len(records) + 1
    formats_seen = {}
    for row, record in zip(ws.iter_rows(min_row=2, max_col=23), records):
        assert len(row) == 23
        price_cell = row[6]
        assert isinstance(price_cell.value, (int, float)) \
            and price_cell.value == record["unit_price"]
        code = price_cell.number_format
        assert code.startswith('"£"0.00')
        suffix = code[len('"£"0.00'):]
        # Excel literal escapes only: no basis character acts as a directive.
        assert re.fullmatch(r"(?:\\[\s\S])*", suffix)
        visible_basis = "".join(re.findall(r"\\([\s\S])", suffix))
        basis = record.get("unit_basis")
        assert visible_basis == (" " + basis.replace("\x00", "") if basis else "")
        formats_seen.setdefault(code, set()).add(basis or None)
        for index, cell in enumerate(row[11:22:2]):
            if index < len(record["overview_options"]):
                assert cell.value == record["overview_options"][index]["price"]
                assert cell.number_format == '"£"0.00'
            else:
                assert cell.value is None
    # identical bases share an identical format string
    for code, seen in formats_seen.items():
        assert len(seen) == 1
    for row, record in zip(details.iter_rows(min_row=2), records):
        assert row[16].number_format.startswith('"£"0.00')
        assert row[19].number_format == "0.00%"
    assert records == before


def test_workbook_raw_record_round_trip_arbitrary_fields_variants_and_dictionary():
    records = [rec("Synthetic long", product_key="synthetic-v1", unit_basis="per g", unit_price=5,
                   raw_options=[{"size": f"{i}g", "price": i * 5,
                                 "variant_id": f"synthetic-option-{i}", "extra": [1, 2]}
                                for i in range(1, 9)],
                   additional_source_metadata={"url": "https://example.invalid/source", "confidence": 0.4},
                   arbitrary_new_field={"text": "=1+1<&>\r\n", "unknown": None},
                   store_description="Synthetic evidence " * 2500),
               rec("Synthetic short", category="Vapes", product_key="synthetic-v2")]
    scored = sort_records(build_factors(records))
    before = copy.deepcopy(records)
    wb = _load(build_workbook(scored, "2026-10-25"))
    ws, details = wb["Overview"], wb["Details"]
    headers = [c.value for c in details[1]]
    part_indices = [i for i, name in enumerate(headers)
                    if isinstance(name, str) and name.startswith("Raw record JSON part ")]
    assert len(part_indices) > 1  # long raw JSON is split, never silently truncated
    assert ws.max_row == details.max_row == len(records) + 1
    for row_number, (record, _, _) in enumerate(scored, start=2):
        raw = "".join(v for i in part_indices
                      if (v := details.cell(row_number, i + 1).value) is not None)
        assert json.loads(raw) == record
        assert ws.cell(row_number, 2).value == record["product_name"] == details.cell(row_number, 3).value
        methodology = json.loads(details.cell(row_number, headers.index("Scoring methodology (JSON)") + 1).value)
        assert methodology["weights"] == {"grade": 1, "rating": 1, "price": 1,
                                          "thc": 1, "si": 1, "popularity": 0.25}
        assert methodology["gamma"] == 6
        assert methodology["canon_ref"]["repository"] == "tyecam1/misumi"
    assert "Field dictionary (JSON)" in headers
    assert records == before


def test_workbook_empty_input_keeps_schema():
    wb = _load(build_workbook([], "2026-10-25"))
    assert wb.sheetnames == ["Overview", "Details"]
    assert [c.value for c in wb["Overview"][1][:23]] == OVERVIEW_COLUMNS
    assert wb["Overview"].max_row == 1
    assert wb["Details"].max_row == 1


# --- combined rating / popularity Overview population (acceptance audit) -----

def test_attach_combined_inputs_blends_raw_ratings():
    records = [
        rec("Rated both", leafly_rating=5.0, leafly_rating_count=2,
            seller_rating=4.0, seller_rating_count=6),
        rec("Leafly only", leafly_rating=4.5, leafly_rating_count=10),
        rec("Unrated zero reviews", seller_rating=0.0, seller_rating_count=0),
        rec("No rating data"),
        rec("Pre-computed", combined_rating=3.14,
            leafly_rating=5.0, leafly_rating_count=99),
        rec("Workbook field name", leafly_rating=5.0, leafly_rating_count=2,
            seller_rating=4.0, seller_review_count=6),
    ]
    from src.product_research import attach_combined_inputs
    attach_combined_inputs(records)
    assert records[0]["combined_rating"] == pytest.approx(
        combined_rating((5.0, 2), (4.0, 6))[0])
    assert records[0]["combined_rating_detail"]["weights"]["leafly"] > 0
    assert records[1]["combined_rating"] == pytest.approx(4.5)
    assert records[2]["combined_rating"] is None  # unrated is not zero stars
    assert "combined_rating" not in records[3]
    assert records[4]["combined_rating"] == pytest.approx(3.14)  # never overwritten
    # seller_review_count (workbook schema name) must carry seller weight too
    assert records[5]["combined_rating"] == pytest.approx(
        combined_rating((5.0, 2), (4.0, 6))[0])


def test_overview_popularity_index_populated_from_factor():
    records = [
        rec("Ranked A", seller_popularity_rank=1, unit_price=10.0, unit_basis="per g"),
        rec("Ranked B", seller_popularity_rank=2, unit_price=10.0, unit_basis="per g"),
    ]
    scored = sort_records(build_factors(records))
    wb = _load(build_workbook(scored, "2026-10-25"))
    ws = wb["Overview"]
    col = OVERVIEW_COLUMNS.index("Combined popularity index / 100") + 1
    values = [ws.cell(row, col).value for row in (2, 3)]
    assert all(v is not None for v in values)  # regression: blank on 5d947fb
    assert all(0.0 <= v <= 100.0 for v in values)
    # better seller rank -> higher index; neutral factor would sit at 50
    assert values[0] > values[1]


def test_overview_combined_rating_populated_after_attach():
    records = [
        rec("Rated A", leafly_rating=5.0, leafly_rating_count=2,
            unit_price=10.0, unit_basis="per g"),
        rec("Rated B", leafly_rating=3.0, leafly_rating_count=8,
            unit_price=10.0, unit_basis="per g"),
    ]
    from src.product_research import attach_combined_inputs
    attach_combined_inputs(records)
    scored = sort_records(build_factors(records))
    wb = _load(build_workbook(scored, "2026-10-25"))
    ws = wb["Overview"]
    col = OVERVIEW_COLUMNS.index("Combined rating / 5") + 1
    values = [ws.cell(row, col).value for row in (2, 3)]
    assert all(v is not None for v in values)  # regression: blank on 5d947fb
    assert all(0.0 <= v <= 5.0 for v in values)


def test_overview_popularity_index_explicit_value_wins():
    records = [
        rec("Explicit", combined_popularity_index=77.0,
            seller_popularity_rank=1, unit_price=10.0, unit_basis="per g"),
        rec("Ranked", seller_popularity_rank=2, unit_price=10.0, unit_basis="per g"),
    ]
    scored = sort_records(build_factors(records))
    wb = _load(build_workbook(scored, "2026-10-25"))
    ws = wb["Overview"]
    col = OVERVIEW_COLUMNS.index("Combined popularity index / 100") + 1
    assert ws.cell(2, col).value == 77.0


def test_combined_rating_identity_confidence_attenuates_source_weight():
    # Same review counts; an uncertain-identity analogue Leafly source must
    # move the blend toward the verified seller value (canon 1.10: never
    # substitute an unsupported Leafly cultivar name as a verified rating).
    exact, d1 = combined_rating((4.0, 50), (2.0, 10))
    analogue, d2 = combined_rating((4.0, 50), (2.0, 10),
                                   leafly_identity_confidence=0.3)
    assert analogue < exact
    assert analogue > 2.0  # still pulled up by the (attenuated) leafly side
    assert d2["leafly_identity_confidence"] == 0.3
    assert d1["leafly_identity_confidence"] == 1.0
    # full confidence reproduces the historical blend exactly
    assert combined_rating((4.0, 50), (2.0, 10))[0] == exact
    # single-source values are unchanged by identity confidence
    solo, _ = combined_rating((4.0, 50), (None, None),
                              leafly_identity_confidence=0.3)
    assert solo == 4.0
    # seller confidence is applied symmetrically: attenuating the seller
    # side reduces its pull, moving the blend toward the Leafly value
    exact_s = combined_rating((4.0, 10), (2.0, 50))[0]
    s_att, _ = combined_rating((4.0, 10), (2.0, 50),
                               seller_identity_confidence=0.3)
    assert s_att > exact_s
    # zero confidence with an unknown count removes the fallback pull
    z, d = combined_rating((4.0, None), (2.0, 10),
                           leafly_identity_confidence=0.0)
    assert z == 2.0 and d["weights"]["leafly"] == 0.0
    # confidence is bounded
    import pytest as _pytest
    with _pytest.raises(ValueError):
        combined_rating((4.0, 10), (2.0, 10), leafly_identity_confidence=1.5)
    with _pytest.raises(ValueError):
        combined_rating((4.0, 10), (2.0, 10), seller_identity_confidence=-0.1)


def test_attach_combined_inputs_passes_identity_confidence():
    from src.product_research import attach_combined_inputs
    records = [rec("Analogue", leafly_rating=5.0, leafly_rating_count=20,
                   seller_rating=1.0, seller_rating_count=20,
                   leafly_identity_confidence=0.3)]
    attach_combined_inputs(records)
    blended, _ = combined_rating((5.0, 20), (1.0, 20),
                                 leafly_identity_confidence=0.3)
    assert records[0]["combined_rating"] == pytest.approx(blended)
    assert records[0]["combined_rating"] < 3.0  # attenuated toward seller
