"""Two-sheet Overview/Details workbook writer per the Misumi canon contract
(exact 23-column Overview A:W, row-aligned Details, neutral category banding).

Repaired in the 2026-10-26 corrective round: built on openpyxl (a supported,
tested spreadsheet implementation) instead of hand-rolled OOXML. Composite
scores are LIVE Excel formulas over Details factor-input cells and editable
configurable weight cells, so editing an input or a weight recalculates every
affected score. The score gradient is a genuine Excel conditional colour
scale, not static per-decile fills. The generation script is the
deterministic regeneration/re-sort mechanism: rows are ordered by
product_research.sort_records before writing, because Excel cannot
automatically re-sort rows when inputs change.
"""

from __future__ import annotations

import json
import math
from datetime import datetime, timezone
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.product_research import (
    CANON_REF, GAMMA, SCORING_METHODOLOGY, W_GRADE, W_POPULARITY, W_PRICE,
    W_RATING, W_SI, W_THC,
)

OVERVIEW_COLUMNS = [
    "Category", "Title", "Green Team description", "Composite score (%)",
    "Combined rating / 5", "Combined popularity index / 100",
    "Minimum price per quantity", "Leafly THC (%)", "Sativa (%)", "Indica (%)",
    "Size 1", "Price 1", "Size 2", "Price 2", "Size 3", "Price 3",
    "Size 4", "Price 4", "Size 5", "Price 5", "Size 6", "Price 6",
    "Buying status (snapshot date)",
]

_BASE_DETAIL_COLUMNS = [
    "Product key", "Category", "Title", "Buying status", "Seller URL",
    "Fetched (UTC)", "Green Team description (full)", "Leafly description",
    "Seller tier (reported)", "Inferred grade (rarity estimate)",
    "Inferred grade confidence", "Leafly rating", "Leafly rating count",
    "Seller rating", "Seller review count", "Seller popularity rank (ordinal)",
    "Unit price (min, GBP)", "Unit basis", "All advertised options (raw)",
    "Composite score (%)", "Normalized factors", "Factor weights used",
    "Missing applicable metrics", "Missing penalty", "Gamma", "GM", "B", "C",
    "Excluded price options", "Match classification", "Primary Leafly source",
    "Factor evidence and peer groups (JSON)", "Scoring methodology (JSON)",
    "Field dictionary (JSON)",
]

# Live-calculation columns appended to Details. The six factor cells are the
# EDITABLE INPUTS of the score formulas; the rest are live formula columns.
FACTOR_INPUT_COLUMNS = [
    "Grade factor (input)", "Rating factor (input)", "Price factor (input)",
    "THC factor (input)", "Sativa/Indica factor (input)",
    "Popularity factor (input)",
]
_CALC_COLUMNS = (
    ["Missing applicable count"]
    + [f"w*ln(f) ({name})" for name in
       ("grade", "rating", "price", "thc", "si", "popularity")]
    + ["Sum of present weights", "GM (live)", "B (live)", "C (live)"]
)
DETAIL_COLUMNS = _BASE_DETAIL_COLUMNS + FACTOR_INPUT_COLUMNS + _CALC_COLUMNS

IDX_MISSING_COUNT = len(_BASE_DETAIL_COLUMNS) + len(FACTOR_INPUT_COLUMNS)
IDX_FACTORS = {name: len(_BASE_DETAIL_COLUMNS) + i
               for i, name in enumerate(FACTOR_INPUT_COLUMNS)}
IDX_CONTRIB = {name: IDX_MISSING_COUNT + 1 + i
               for i, name in
               enumerate(("grade", "rating", "price", "thc", "si", "popularity"))}
IDX_WSUM = IDX_CONTRIB["popularity"] + 1
IDX_GM = IDX_WSUM + 1
IDX_B = IDX_GM + 1
IDX_C = IDX_B + 1

# Configurable weight cells live on Overview row 1 beyond the fixed 23-column
# table (Y..AF), where they can never collide with data or raw-JSON parts.
WEIGHT_SHEET = "Overview"
WEIGHT_CELLS = {
    "grade": "Z1", "rating": "AA1", "price": "AB1",
    "thc": "AC1", "si": "AD1", "popularity": "AE1", "gamma": "AF1",
}
WEIGHT_DEFAULTS = {"grade": W_GRADE, "rating": W_RATING, "price": W_PRICE,
                   "thc": W_THC, "si": W_SI, "popularity": W_POPULARITY}
_WEIGHT_TITLE = ("Configurable weights, order grade,rating,price,thc,si,"
                 "popularity (edit Z1:AE1, gamma in AF1; all scores "
                 "recalculate; the per-row factor (input) cells in Details "
                 "feed the formulas. Regeneration/re-sort: rerun the "
                 "deterministic build; Excel does not re-sort automatically)")
_WEIGHT_LABELS = [("Y1", _WEIGHT_TITLE)]


def _abs_ref(ref: str) -> str:
    """'BH1' -> '$BH$1' (works for one- and two-letter columns)."""
    return f"${ref[:-1]}${ref[-1]}"

FIELD_DICTIONARY = {
    "Composite score (%)": "Relative comparison, not quality certainty; stored on a 0..1 scale; a LIVE formula over the Details factor inputs and weights",
    "Leafly THC (%)": "Cultivar reference THC, not a batch measurement; source percentage points",
    "Sativa (%) / Indica (%)": "Stated source orientation in percentage points, not measured effects",
    "Unit price": "Minimum defensible price per stated unit; compare category AND unit basis only",
    "Inferred grade": "Injected estimate with confidence/provenance, distinct from seller-reported tier",
    "Normalized factors": "Category-relative multipliers; 1 is neutral; absent is not zero",
    "Factor inputs (live calculation)": "The six '<metric> factor (input)' cells feed the live score formulas; edit one and its row's GM/B/C/score recalculate",
    "Regeneration and re-sort": "Rerun the deterministic build to re-sort: row order (availability, category, score desc, unit price asc) is fixed at generation; Excel does not re-sort automatically",
    "Factor evidence and peer groups (JSON)": "Raw factor inputs, confidence, peer vectors and source pointers",
    "Raw record JSON parts": "Concatenate numbered parts in this row, then decode JSON for ALL input fields, variants and source metadata",
    "Text exceeding a cell limit": "Display preview is labelled; full value retained in raw record JSON parts",
}

JSON_CHUNK_SIZE = 30000  # below Excel's 32767-character cell text limit
CELL_TEXT_LIMIT = 32767
MONEY_FMT = '"£"0.00'
SCORE_FMT = "0.00%"
FACTOR_FMT = "0.0000"
SOURCE_PERCENT_FMT = '0.0"%"'


def _json(value: Any) -> str:
    # ASCII escapes also preserve control characters that XML cannot represent.
    return json.dumps(value, ensure_ascii=True, sort_keys=True)


def _clean(text: Any) -> str:
    """Strip XML-invalid control characters; label previews above the limit."""
    s = ILLEGAL_CHARACTERS_RE.sub("", str(text))
    if len(s) > CELL_TEXT_LIMIT:
        s = s[:JSON_CHUNK_SIZE] + " [preview; full value in Details raw record JSON]"
    return s


def _factor_values(factors) -> Dict[str, Any]:
    """The six editable factor inputs for one record."""
    return {
        "grade": factors.grade[0] if factors.grade else None,
        "rating": factors.rating,
        "price": factors.price[0] if factors.price else None,
        "thc": factors.thc,
        "si": factors.si,
        "popularity": factors.popularity,
    }


def build_workbook(scored_sorted, snapshot_date: str) -> bytes:
    """Return XLSX bytes; never opens or overwrites a user's workbook.

    scored_sorted is the output of product_research.sort_records. All records
    must be JSON-serializable; unsupported values fail explicitly.
    """
    scored_sorted = list(scored_sorted)
    raw_records = [_json(rec) for rec, _, _ in scored_sorted]
    raw_parts = [[raw[i:i + JSON_CHUNK_SIZE] for i in range(0, len(raw), JSON_CHUNK_SIZE)]
                 for raw in raw_records]
    part_count = max((len(parts) for parts in raw_parts), default=1)
    detail_headers = DETAIL_COLUMNS + [f"Raw record JSON part {i}"
                                       for i in range(1, part_count + 1)]

    wb = Workbook()
    wb.properties.creator = "odysseus"
    stamp = datetime(2026, 10, 26, tzinfo=timezone.utc)
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.calculation.fullCalcOnLoad = True  # Excel/LibreOffice recalculate on open
    overview = wb.active
    overview.title = "Overview"
    details = wb.create_sheet("Details")
    top_align = Alignment(wrap_text=True, vertical="top")
    hdr_font = Font(bold=True)

    categories: List[str] = []
    for rec, _f, _s in scored_sorted:
        if rec["category"] not in categories:
            categories.append(rec["category"])
    band_fill = {}
    for i, cat in enumerate(categories):
        fill = PatternFill("solid", fgColor="F2F2F2" if i % 2 else "FFFFFF")
        band_fill[cat] = fill
    money_basis_fmt: Dict[str, str] = {}

    def basis_number_format(basis) -> str:
        if not basis:
            return MONEY_FMT
        basis = "".join(ch for ch in str(basis) if not ILLEGAL_CHARACTERS_RE.match(ch))
        if basis not in money_basis_fmt:
            # Excel's backslash quotes the next character literally, so no
            # basis character can act as a format directive.
            money_basis_fmt[basis] = MONEY_FMT + "".join("\\" + ch
                                                         for ch in " " + basis)
        return money_basis_fmt[basis]

    # -- headers -----------------------------------------------------------
    for col, name in enumerate(OVERVIEW_COLUMNS, start=1):
        cell = overview.cell(1, col, name)
        cell.font = hdr_font
        cell.alignment = top_align
    for col, name in enumerate(detail_headers, start=1):
        cell = details.cell(1, col, name)
        cell.font = hdr_font
        cell.alignment = top_align
    for ref, label in _WEIGHT_LABELS:
        overview[ref] = label
        overview[ref].alignment = top_align
    for metric, ref in WEIGHT_CELLS.items():
        overview[ref] = (GAMMA if metric == "gamma"
                         else WEIGHT_DEFAULTS[metric])

    # -- rows ----------------------------------------------------------------
    for r, (rec, factors, score) in enumerate(scored_sorted, start=2):
        opts = rec.get("overview_options") or []
        band = band_fill[rec["category"]]
        ovals = [
            rec["category"], rec["product_name"],
            rec.get("store_description") or rec.get("product_name", ""),
            None,  # live score formula below
            rec.get("combined_rating"), rec.get("combined_popularity_index"),
            rec.get("unit_price"), rec.get("leafly_thc"), rec.get("sativa_pct"),
            rec.get("indica_pct"),
        ]
        for k in range(6):
            ovals += [opts[k].get("size") if k < len(opts) else None,
                      opts[k].get("price") if k < len(opts) else None]
        ovals.append(f"{rec.get('buying_status', 'unknown')} (snapshot {snapshot_date})")
        for col, val in enumerate(ovals, start=1):
            cell = overview.cell(r, col)
            if col == 4:
                cell.value = (
                    f"=MIN(1,MAX(0,Details!{get_column_letter(IDX_C + 1)}{r}"
                    f"*0.8^MAX(0,Details!{get_column_letter(IDX_MISSING_COUNT + 1)}{r}-1)))")
                cell.number_format = SCORE_FMT
            elif isinstance(val, (int, float)) and not isinstance(val, bool) \
                    and math.isfinite(val):
                cell.value = val
                if col == 7:
                    cell.number_format = basis_number_format(rec.get("unit_basis"))
                elif col in (8, 9, 10):
                    cell.number_format = SOURCE_PERCENT_FMT
                elif col >= 12 and col % 2 == 0:
                    cell.number_format = MONEY_FMT
            elif val is not None:
                cell.value = _clean(val)
            if col != 4:
                cell.fill = band
                cell.alignment = top_align

        # Details live-calculation chain
        fv = _factor_values(factors)
        metric_order = ("grade", "rating", "price", "thc", "si", "popularity")
        for name in metric_order:
            cell = details.cell(r, IDX_FACTORS[_FACTOR_COLUMN_NAME[name]] + 1)
            cell.value = fv[name]
            cell.number_format = FACTOR_FMT
            cell.alignment = top_align
        first_factor_letter = get_column_letter(IDX_FACTORS[FACTOR_INPUT_COLUMNS[0]] + 1)
        weight_range = f"{WEIGHT_SHEET}!Z$1:AE$1"
        for name in metric_order:
            fletter = get_column_letter(
                IDX_FACTORS[_FACTOR_COLUMN_NAME[name]] + 1)
            contrib = details.cell(r, IDX_CONTRIB[name] + 1)
            wref = WEIGHT_CELLS[name]
            contrib.value = (f'=IF({fletter}{r}="",0,'
                             f'{WEIGHT_SHEET}!{_abs_ref(wref)}*LN({fletter}{r}))')
            contrib.number_format = FACTOR_FMT
            contrib.alignment = top_align
        wsum = details.cell(r, IDX_WSUM + 1)
        wsum.value = (f"=SUMPRODUCT(({first_factor_letter}{r}:"
                      f"{get_column_letter(IDX_FACTORS[FACTOR_INPUT_COLUMNS[-1]] + 1)}{r}"
                      f'<>"")*{weight_range})')
        wsum.number_format = FACTOR_FMT
        c_letters = [get_column_letter(IDX_CONTRIB[n] + 1) for n in metric_order]
        gm = details.cell(r, IDX_GM + 1)
        gm.value = (f"=IF({get_column_letter(IDX_WSUM + 1)}{r}<=0,1,"
                    f"EXP(({c_letters[0]}{r}+{c_letters[1]}{r}+{c_letters[2]}{r}"
                    f"+{c_letters[3]}{r}+{c_letters[4]}{r}+{c_letters[5]}{r})"
                    f"/{get_column_letter(IDX_WSUM + 1)}{r}))")
        gm.number_format = FACTOR_FMT
        b_cell = details.cell(r, IDX_B + 1)
        gm_letter = get_column_letter(IDX_GM + 1)
        ws_letter = get_column_letter(IDX_WSUM + 1)
        b_cell.value = (f"=MIN(1,MAX(0,0.5+({gm_letter}{r}-1)*{ws_letter}{r}"
                        f"/({ws_letter}{r}+2)))")
        b_cell.number_format = FACTOR_FMT
        b_letter = get_column_letter(IDX_B + 1)
        c_cell = details.cell(r, IDX_C + 1)
        gamma_ref = f"{WEIGHT_SHEET}!{_abs_ref(WEIGHT_CELLS['gamma'])}"
        c_cell.value = (f"=IF(AND({b_letter}{r}>0,{b_letter}{r}<1),"
                        f"{b_letter}{r}^{gamma_ref}"
                        f"/({b_letter}{r}^{gamma_ref}"
                        f"+(1-{b_letter}{r})^{gamma_ref}),"
                        f"{b_letter}{r})")
        c_cell.number_format = FACTOR_FMT
        m_cell = details.cell(r, IDX_MISSING_COUNT + 1)
        m_cell.value = len(score["missing_applicable"])

        # static audit columns
        dvals = [
            rec.get("product_key") or rec["product_name"], rec["category"],
            rec["product_name"], rec.get("buying_status", "unknown"),
            rec.get("permalink") or "", rec.get("fetched") or "",
            rec.get("store_description") or "", rec.get("leafly_description") or "",
            rec.get("seller_tier") or "", (rec.get("inferred_grade") or {}).get("tier", ""),
            (rec.get("inferred_grade") or {}).get("confidence"),
            rec.get("leafly_rating"), rec.get("leafly_rating_count"),
            rec.get("seller_rating"), rec.get("seller_review_count"),
            rec.get("seller_popularity_rank"),
            rec.get("unit_price"), rec.get("unit_basis"),
            _json(rec.get("raw_options")), score["score"],
            _json({k: v for k, v in dict(grade=factors.grade, rating=factors.rating,
                                         price=factors.price, thc=factors.thc,
                                         si=factors.si,
                                         popularity=factors.popularity).items()
                   if v is not None}),
            score["weights_used"], "; ".join(score["missing_applicable"]),
            score["missing_penalty"], score["gamma"], score["gm"], score["b"],
            score["c"],
            _json(rec.get("price_detail", {}).get("excluded_options", [])),
            rec.get("match_classification") or "",
            rec.get("primary_leafly_source") or "",
            _json(factors.detail),
            _json({"canon_ref": CANON_REF, **SCORING_METHODOLOGY}),
            _json(FIELD_DICTIONARY),
        ]
        dvals.extend(raw_parts[r - 2] + [None] * (part_count - len(raw_parts[r - 2])))
        part_start = len(DETAIL_COLUMNS) + 1  # raw parts live beyond the calc block
        for col, val in enumerate(dvals, start=1):
            if col > len(_BASE_DETAIL_COLUMNS):
                continue  # factor/calc columns were written as live cells above
            cell = details.cell(r, col)
            if isinstance(val, (int, float)) and not isinstance(val, bool) \
                    and math.isfinite(val):
                cell.value = val
            elif val not in (None, ""):
                cell.value = _clean(val)
            if col == 17:
                cell.number_format = basis_number_format(rec.get("unit_basis"))
            elif col == 20:
                cell.number_format = SCORE_FMT
            cell.alignment = top_align
        for i, part in enumerate(raw_parts[r - 2]):
            cell = details.cell(r, part_start + i, _clean(part))
            cell.alignment = top_align

    # genuine Excel conditional colour scale on the score column only
    if overview.max_row >= 2:
        overview.conditional_formatting.add(
            f"D2:D{overview.max_row}",
            ColorScaleRule(
                start_type="num", start_value=0, start_color="E8F5E9",
                mid_type="num", mid_value=0.5, mid_color="AED9B0",
                end_type="num", end_value=1, end_color="4C9A52"))

    overview.freeze_panes = "A2"
    details.freeze_panes = "A2"
    widths_o = [24, 30, 60, 20, 20, 24, 24, 18, 16, 16] + [14, 14] * 6 + [42]
    for i, w in enumerate(widths_o, start=1):
        overview.column_dimensions[get_column_letter(i)].width = w
    for i in range(1, len(detail_headers) + 1):
        details.column_dimensions[get_column_letter(i)].width = \
            60 if i in (7, 8) else 28
    for letter in ("Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF"):
        overview.column_dimensions[letter].width = 30

    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_FACTOR_COLUMN_NAME = {
    "grade": "Grade factor (input)",
    "rating": "Rating factor (input)",
    "price": "Price factor (input)",
    "thc": "THC factor (input)",
    "si": "Sativa/Indica factor (input)",
    "popularity": "Popularity factor (input)",
}
